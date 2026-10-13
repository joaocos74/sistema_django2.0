from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from ..models import Estabelecimento, Tipo, Responsavel, RTResponsavel, Vigilancia, Documento
from denuncias.models import Denuncia
from redesim.models import RegistroRedesim
from termos.models import AutoTermoPadrao, NotificacaoPadrao, NotificacaoEmitida


class ConsolidadoExcelTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('fiscal')
        self.client.force_login(self.user)
        self.url = reverse('register_view:exportar_consolidado')

    def workbook(self):
        response = self.client.get(self.url, {'ano': '1900', 'fiscal': '999'})
        self.assertEqual(response.status_code, 200)
        dados = b''.join(response.streaming_content)
        with patch('django.db.backends.base.base.BaseDatabaseWrapper.close_if_unusable_or_obsolete'):
            response.close()
        return load_workbook(BytesIO(dados))

    def test_todos_vinculos_historicos_e_textos(self):
        est = Estabelecimento.objects.create(nome_fantasia='=1+1', cnpj='001234', status=False)
        Estabelecimento.objects.create(nome_fantasia='Inativo', status=True)
        for classe in ('Antiga', 'Atual'):
            Tipo.objects.create(estabelecimento=est, classe=classe)
        Responsavel.objects.create(estabelecimento=est, nome='Responsável', cpf='001')
        RTResponsavel.objects.create(estabelecimento=est, nome='Técnico')
        Vigilancia.objects.create(estabelecimento=est, fiscal=self.user)
        RegistroRedesim.objects.create(estabelecimento=est, fiscal=self.user)
        Denuncia.objects.create(estabelecimento=est, criado_por=self.user, motivo='x' * 40000)
        Denuncia.objects.create(criado_por=self.user)
        Documento.objects.create(estabelecimento=est, nome='Anexo', arquivo='documentos/teste.pdf', enviado_por=self.user)
        AutoTermoPadrao.objects.create(estabelecimento=est, atualizado_por=self.user, campos={'campo_37': '=texto'})
        NotificacaoPadrao.objects.create(estabelecimento=est, atualizado_por=self.user, campos={'campo_33': 'Padrão'})
        NotificacaoEmitida.objects.create(estabelecimento=est, gerado_por=self.user, numero=1, ano=2026,
            chave=uuid4(), campos={'campo_33': 'Snapshot', 'autoridades_extras': [{'nome': 'Fiscal'}]})
        wb = self.workbook()
        self.assertEqual(wb.sheetnames[0], 'Consolidado por estabelecimento')
        principal = wb.worksheets[0]
        headers = [cell.value for cell in principal[1]]
        linhas = list(principal.values)[1:]
        vinculadas = [row for row in linhas if row[0] == est.pk]
        self.assertEqual(len(vinculadas), 2)  # O histórico extenso continua sem truncamento.
        classificacoes = vinculadas[0][headers.index('Classificações')]
        self.assertIn('Antiga', classificacoes)
        self.assertIn('Atual', classificacoes)
        self.assertIn('Snapshot', vinculadas[0][headers.index('Notificações geradas')])
        self.assertIn('Responsável', vinculadas[0][headers.index('Responsáveis')])
        self.assertEqual(len(linhas), 3)  # Inclui também o cadastro inativo sem vínculos.
        self.assertEqual(wb['Estabelecimentos'].max_row, 3)
        self.assertEqual(wb['Classificações'].max_row, 3)
        self.assertEqual(wb['Denúncias'].max_row, 3)
        for nome in ['Responsáveis', 'Responsáveis técnicos', 'Inspeções', 'REDESIM', 'Documentos',
                     'Padrões Auto Termo', 'Padrões Notificação', 'Notificações geradas']:
            self.assertEqual(wb[nome].max_row, 2, nome)
        cabecalho = [c.value for c in wb['Estabelecimentos'][1]]
        cell = wb['Estabelecimentos'].cell(2, cabecalho.index('nome_fantasia') + 1)
        self.assertEqual(cell.value, '=1+1')
        self.assertEqual(cell.data_type, 's')
        self.assertEqual(wb['Estabelecimentos'].cell(2, cabecalho.index('cnpj') + 1).value, '001234')
        self.assertEqual(''.join(row[-1] for row in list(wb['Textos extensos'].values)[1:]), 'x' * 40000)
        self.assertIn('Snapshot', [row[-1] for row in wb['Conteúdo termos'].values])
        self.assertIn('/documentos/', wb['Documentos'].cell(2, wb['Documentos'].max_column).value)
        self.assertFalse(any('password' in str(c.value) for ws in wb for c in ws[1]))
        wb.close()

    def test_vazio_e_acesso(self):
        wb = self.workbook()
        self.assertEqual(wb['Estabelecimentos'].max_row, 1)
        wb.close()
        self.assertEqual(self.client.post(self.url).status_code, 405)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
