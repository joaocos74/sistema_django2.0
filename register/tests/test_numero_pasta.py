from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import connection, connections, IntegrityError, transaction
from django.db.migrations.executor import MigrationExecutor
from django.test import Client, RequestFactory, TestCase, TransactionTestCase
from django.urls import reverse
from openpyxl import load_workbook

from denuncias.models import Denuncia
from redesim.models import RegistroRedesim
from termos.models import AutoTermoPadrao, NotificacaoEmitida
from ..exportacao_excel import gerar_excel
from ..models import Estabelecimento, HistoricoPasta, SequenciaPasta, Vigilancia


class NumeroPastaTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser('admin', password='teste')
        self.fiscal = get_user_model().objects.create_user('fiscal')
        self.est = Estabelecimento.objects.create(nome_fantasia='Mercado A', cnpj='123', numero_pasta=25)
        self.vigilancia = Vigilancia.objects.create(estabelecimento=self.est, fiscal=self.fiscal,
            ano_previsto=2026, mes_previsto=10)
        self.url = reverse('register_view:alterar_numero_pasta', args=[self.est.pk])
        self.client.force_login(self.admin)

    def test_sequencia_considera_ativos_e_reaproveita_inativo(self):
        self.est.numero_pasta = 40
        self.est.save(usuario=self.admin)
        self.est.numero_pasta = 25
        self.est.status = True
        self.est.save(usuario=self.admin)
        outro = Estabelecimento.objects.create(nome_fantasia='Mercado B', numero_pasta=25)
        novo = Estabelecimento.objects.create(nome_fantasia='Mercado C')
        self.assertEqual(novo.numero_pasta, 26)
        self.assertNotEqual(outro.pk, self.est.pk)
        self.vigilancia.refresh_from_db()
        self.assertEqual(self.vigilancia.estabelecimento_id, self.est.pk)
        self.assertFalse(outro.fiscalizacoes.exists())
        self.assertEqual(self.est.historico_pastas.first().numero_anterior, 40)
        self.assertEqual(self.est.historico_pastas.first().alterado_por, self.admin)
        self.est.status = False
        with self.assertRaisesMessage(ValidationError, 'já pertence'):
            self.est.save()
        self.est.refresh_from_db()
        self.assertTrue(self.est.status)
        self.est.numero_pasta = 42
        self.est.status = False
        self.est.save()

    def test_inativar_um_e_trocar_seis_por_um_proximo_recebe_seis(self):
        self.est.status = True
        self.est.save()
        cadastros = [Estabelecimento.objects.create(nome_fantasia=f'Teste {n}') for n in range(1, 7)]
        self.assertEqual([est.numero_pasta for est in cadastros], [1, 2, 3, 4, 5, 6])
        cadastros[0].status = True
        cadastros[0].save()
        cadastros[-1].numero_pasta = 1
        cadastros[-1].save()
        novo = Estabelecimento.objects.create(nome_fantasia='Novo')
        self.assertEqual(novo.numero_pasta, 6)
        self.assertNotIn(novo.pk, [est.pk for est in cadastros])

    def test_sem_ativos_comeca_em_um(self):
        self.est.status = True
        self.est.save()
        self.assertEqual(Estabelecimento.objects.create(nome_fantasia='Novo').numero_pasta, 1)

    def test_lacuna_intermediaria_nao_muda_maior_ativo(self):
        Estabelecimento.objects.create(nome_fantasia='Pasta menor', numero_pasta=2)
        Estabelecimento.objects.create(nome_fantasia='Inativo maior', numero_pasta=500, status=True)
        self.assertEqual(Estabelecimento.objects.create(nome_fantasia='Novo').numero_pasta, 26)

    def test_duplicidade_bloqueada_tambem_no_banco(self):
        outro = Estabelecimento.objects.create(nome_fantasia='Outro')
        with self.assertRaises(IntegrityError), transaction.atomic():
            Estabelecimento.objects.filter(pk=outro.pk).update(numero_pasta=25)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Estabelecimento.objects.filter(pk=outro.pk).update(numero_pasta=0)

    def test_falha_reverte_numeracao_e_historico(self):
        with self.assertRaises(ValueError), transaction.atomic():
            self.est.numero_pasta = 90
            self.est.save(usuario=self.admin)
            raise ValueError('Falha posterior')
        self.est.refresh_from_db()
        self.assertEqual(self.est.numero_pasta, 25)
        self.assertEqual(self.est.historico_pastas.count(), 1)
        self.assertEqual(Estabelecimento.objects.create(nome_fantasia='Novo').numero_pasta, 26)

    def test_admin_altera_sem_criar_inspecao_e_fiscal_nao_altera(self):
        self.client.force_login(self.fiscal)
        self.assertEqual(self.client.post(self.url, {'numero_pasta': 50}).status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(self.url).status_code, 405)
        self.assertEqual(self.client.post(self.url, {'numero_pasta': 50}).status_code, 302)
        self.est.refresh_from_db()
        self.assertEqual(self.est.numero_pasta, 50)
        self.assertEqual(self.est.fiscalizacoes.count(), 1)
        self.assertEqual(self.est.fiscalizacoes.get().fiscal, self.fiscal)
        self.assertEqual(self.est.historico_pastas.first().alterado_por, self.admin)
        outro = Estabelecimento.objects.create(nome_fantasia='Ocupada', numero_pasta=60)
        response = self.client.post(self.url, {'numero_pasta': outro.numero_pasta}, follow=True)
        self.assertContains(response, 'já pertence')
        self.est.refresh_from_db()
        self.assertEqual(self.est.numero_pasta, 50)
        self.assertEqual(self.client.post(self.url, {'numero_pasta': -1}).status_code, 302)
        seguro = Client(enforce_csrf_checks=True)
        seguro.force_login(self.admin)
        self.assertEqual(seguro.post(self.url, {'numero_pasta': 70}).status_code, 403)

    def test_pesquisa_ordena_filtra_pasta_e_oculta_inativos(self):
        outro = Estabelecimento.objects.create(nome_fantasia='Primeiro', numero_pasta=2)
        inativo = Estabelecimento.objects.create(nome_fantasia='Inativo', numero_pasta=1, status=True)
        url = reverse('register_view:pesquisar')
        self.assertEqual(list(self.client.get(url).context['estabelecimentos']), [outro, self.est])
        self.assertEqual(list(self.client.get(url, {'id': 25}).context['estabelecimentos']), [self.est])
        self.assertEqual(list(self.client.get(url, {'ativo': '0'}).context['estabelecimentos']), [inativo])
        self.assertEqual(len(self.client.get(url, {'ativo': ''}).context['estabelecimentos']), 3)
        response = self.client.get(reverse('register_view:cronograma'))
        self.assertNotContains(response, 'Inativo')
        self.assertContains(response, '<td>25</td>', html=True)

    def test_relatorios_e_apis_preservam_vinculos_apos_renumerar(self):
        denuncia = Denuncia.objects.create(estabelecimento=self.est, criado_por=self.fiscal)
        RegistroRedesim.objects.create(estabelecimento=self.est, fiscal=self.fiscal, data_liberacao=date(2026, 10, 1))
        emitida = NotificacaoEmitida.objects.create(estabelecimento=self.est, numero=1, ano=2026,
            gerado_por=self.fiscal, chave=uuid4(), campos={'campo_04': 'Nome original'})
        AutoTermoPadrao.objects.create(estabelecimento=self.est, atualizado_por=self.fiscal, campos={'campo_03': '25'})
        stats_url = reverse('register_view:estatisticas')
        antes = self.client.get(stats_url, {'ano': '2026', 'fiscal': self.fiscal.pk}).context
        self.client.post(self.url, {'numero_pasta': 99})
        depois = self.client.get(stats_url, {'ano': '2026', 'fiscal': self.fiscal.pk}).context
        for chave in ('metricas_ano', 'metricas_risco', 'metricas_quadrimestre', 'metricas_meses', 'metricas_classe', 'metricas_pendencias', 'metricas_modulos'):
            self.assertEqual(antes[chave], depois[chave], chave)
        denuncia.refresh_from_db()
        emitida.refresh_from_db()
        self.assertEqual(denuncia.estabelecimento_id, self.est.pk)
        self.assertEqual(emitida.campos, {'campo_04': 'Nome original'})
        response = self.client.get(reverse('termos:auto_termo', args=[self.est.pk]))
        self.assertEqual(response.context['campo_03'], 99)
        self.assertNotIn('campo_03', response.context['padrao_salvo'])
        dados = self.client.get(reverse('redesim:buscar_cnpj'), {'documento': '123'}).json()['dados']
        self.assertEqual((dados['id'], dados['numero_pasta']), (self.est.pk, 99))
        dados = self.client.get(reverse('denuncias:buscar_cnpj'), {'cnpj': '123'}).json()['dados']
        self.assertEqual((dados['id'], dados['numero_pasta']), (self.est.pk, 99))
        self.assertContains(self.client.get(reverse('denuncias:index')), '>99</a>')
        with BytesIO() as arquivo:
            gerar_excel(RequestFactory().get('/'), arquivo)
            arquivo.seek(0)
            wb = load_workbook(arquivo)
            self.assertEqual(wb.worksheets[0].cell(2, 1).value, 99)
            for nome in ('Inspeções', 'Denúncias', 'REDESIM', 'Notificações geradas'):
                cabecalhos = [c.value for c in wb[nome][1]]
                self.assertEqual(wb[nome].cell(2, cabecalhos.index('Número da pasta') + 1).value, 99)
                self.assertEqual(wb[nome].cell(2, cabecalhos.index('estabelecimento_id') + 1).value, self.est.pk)
            wb.close()


@skipUnless(connection.vendor == 'postgresql', 'Concorrência validada no PostgreSQL de produção.')
class ConcorrenciaPastaTests(TransactionTestCase):
    def test_cadastros_simultaneos_recebem_numeros_diferentes(self):
        SequenciaPasta.objects.update_or_create(pk=1, defaults={'ultimo_numero': 100})
        barreira = Barrier(2)

        def criar(indice):
            try:
                barreira.wait(timeout=10)
                return Estabelecimento.objects.create(nome_fantasia=f'Local {indice}').numero_pasta
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            resultados = list(pool.map(criar, (1, 2)))
        self.assertEqual(sorted(resultados), [1, 2])

    def test_duas_alteracoes_nao_ocupam_mesma_pasta(self):
        primeiro = Estabelecimento.objects.create(nome_fantasia='Primeiro')
        segundo = Estabelecimento.objects.create(nome_fantasia='Segundo')
        barreira = Barrier(2)

        def alterar(pk):
            try:
                est = Estabelecimento.objects.get(pk=pk)
                est.numero_pasta = 99
                barreira.wait(timeout=10)
                try:
                    est.save()
                    return True
                except ValidationError:
                    return False
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            resultados = list(pool.map(alterar, (primeiro.pk, segundo.pk)))
        self.assertEqual(sorted(resultados), [False, True])
        self.assertEqual(Estabelecimento.objects.filter(numero_pasta=99).count(), 1)


class MigracaoPastaTests(TransactionTestCase):
    def test_preserva_ids_vinculos_e_inicializa_sequencia(self):
        executor = MigrationExecutor(connection)
        finais = executor.loader.graph.leaf_nodes()
        anterior = [('register', '0009_remove_notificacaopadrao_atualizado_por_and_more')]
        try:
            executor.migrate(anterior)
            apps = executor.loader.project_state(anterior).apps
            Est = apps.get_model('register', 'Estabelecimento')
            Fiscal = apps.get_model('auth', 'User')
            Vig = apps.get_model('register', 'Vigilancia')
            fiscal = Fiscal.objects.create(username='migracao')
            est = Est.objects.create(id=400, nome_fantasia='Existente')
            vig = Vig.objects.create(estabelecimento=est, fiscal=fiscal)
            executor = MigrationExecutor(connection)
            executor.migrate(finais)
            self.assertEqual(Estabelecimento.objects.get(pk=400).numero_pasta, 400)
            self.assertEqual(Vigilancia.objects.get(pk=vig.pk).estabelecimento_id, 400)
            self.assertEqual(HistoricoPasta.objects.get(estabelecimento_id=400).numero_novo, 400)
            self.assertEqual(Estabelecimento.objects.create(nome_fantasia='Novo').numero_pasta, 401)
        finally:
            MigrationExecutor(connection).migrate(finais)
