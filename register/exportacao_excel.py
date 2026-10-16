"""Consolidado relacional completo, sem multiplicar linhas por cruzamento de históricos."""
import json
import re
from datetime import datetime
from tempfile import SpooledTemporaryFile

from django.contrib.auth.decorators import login_required
from django.db import models, transaction
from django.http import FileResponse
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from denuncias.models import Denuncia
from redesim.models import RegistroRedesim
from termos.models import AutoTermoPadrao, NotificacaoPadrao, NotificacaoEmitida
from .models import Estabelecimento, Tipo, Responsavel, RTResponsavel, Vigilancia, Documento, HistoricoPasta


ABAS = [
    ('Estabelecimentos', Estabelecimento), ('Classificações', Tipo),
    ('Responsáveis', Responsavel), ('Responsáveis técnicos', RTResponsavel),
    ('Inspeções', Vigilancia), ('REDESIM', RegistroRedesim), ('Denúncias', Denuncia),
    ('Documentos', Documento), ('Padrões Auto Termo', AutoTermoPadrao),
    ('Padrões Notificação', NotificacaoPadrao), ('Notificações geradas', NotificacaoEmitida),
    ('Histórico de pastas', HistoricoPasta),
]
INVALIDOS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def linha(aba, valores, cabecalho=False):
    celulas = []
    for valor in valores:
        if isinstance(valor, datetime) and timezone.is_aware(valor):
            valor = timezone.localtime(valor).replace(tzinfo=None)
        if isinstance(valor, str):
            valor = INVALIDOS.sub('', valor)
        celula = WriteOnlyCell(aba, value=valor)
        celula.alignment = Alignment(vertical='top', wrap_text=True)
        if isinstance(valor, str):
            # Nomes, observações e identificadores são texto, nunca fórmulas Excel.
            celula.data_type = 's'
        if isinstance(valor, datetime):
            celula.number_format = 'dd/mm/yyyy hh:mm:ss'
        elif hasattr(valor, 'year'):
            celula.number_format = 'dd/mm/yyyy'
        if cabecalho:
            celula.font = Font(bold=True, color='FFFFFF')
            celula.fill = PatternFill('solid', fgColor='169B5B')
        celulas.append(celula)
    aba.append(celulas)


def nova_aba(wb, nome, titulos):
    aba = wb.create_sheet(nome)
    aba.freeze_panes = 'A2'
    for i in range(1, len(titulos) + 1):
        aba.column_dimensions[get_column_letter(i)].width = 25
    linha(aba, titulos, True)
    return aba


def descricao_registro(registro, request):
    partes = []
    for campo in registro._meta.concrete_fields:
        if campo.name == 'estabelecimento':
            continue
        valor = getattr(registro, campo.name)
        if campo.is_relation:
            valor = f'{valor.pk} — {valor.get_full_name() or valor.username}' if valor else ''
        elif isinstance(campo, models.JSONField):
            valor = json.dumps(valor, ensure_ascii=False, indent=2)
        elif campo.choices:
            valor = getattr(registro, f'get_{campo.name}_display')()
        elif isinstance(valor, datetime):
            valor = timezone.localtime(valor).strftime('%d/%m/%Y %H:%M:%S') if timezone.is_aware(valor) else str(valor)
        partes.append(f'{campo.verbose_name}: {valor if valor is not None else ""}')
    if isinstance(registro, Documento):
        partes.append('Download: ' + request.build_absolute_uri(reverse('register_view:baixar_documento', args=[registro.pk])))
    if isinstance(registro, NotificacaoEmitida):
        partes.append('Número: ' + registro.numero_formatado)
        partes.append('Documento: ' + request.build_absolute_uri(reverse('termos:emitida', args=[registro.pk])))
    return '\n'.join(partes)


def consolidado_por_estabelecimento(wb, request):
    campos = [Estabelecimento._meta.get_field('numero_pasta')] + [
        c for c in Estabelecimento._meta.concrete_fields if c.name != 'numero_pasta']
    titulos = [('ID interno (referência)' if c.primary_key else str(c.verbose_name)) for c in campos] + ['Parte'] + [nome for nome, _ in ABAS[1:]]
    aba = nova_aba(wb, 'Consolidado por estabelecimento', titulos)
    aba.freeze_panes = 'D2'
    relacoes = []
    for _, modelo in ABAS[1:]:
        usuarios = [f.name for f in modelo._meta.concrete_fields if f.is_relation and f.name != 'estabelecimento']
        relacoes.append(models.Prefetch(modelo._meta.get_field('estabelecimento').remote_field.related_name,
            queryset=modelo.objects.select_related(*usuarios).order_by('pk')))
    for i in range(len(campos) + 2, len(titulos) + 1):
        aba.column_dimensions[get_column_letter(i)].width = 55
    total = 0
    for est in Estabelecimento.objects.order_by('numero_pasta', 'pk').prefetch_related(*relacoes).iterator(chunk_size=100):
        blocos = []
        for _, modelo in ABAS[1:]:
            campo = modelo._meta.get_field('estabelecimento')
            vinculo = getattr(est, campo.remote_field.related_name, None)
            registros = ([vinculo] if vinculo else []) if campo.one_to_one else vinculo.all()
            texto = '\n\n──────────\n\n'.join(descricao_registro(r, request) for r in registros)
            blocos.append([texto[inicio:inicio + 32000] for inicio in range(0, len(texto), 32000)] or [''])
        partes = max(len(bloco) for bloco in blocos)
        for indice in range(partes):
            valores = [getattr(est, campo.name) for campo in campos]
            valores += [f'{indice + 1}/{partes}'] + [bloco[indice] if indice < len(bloco) else '' for bloco in blocos]
            linha(aba, valores)
            total += 1
    aba.auto_filter.ref = f'A1:{get_column_letter(len(titulos))}{total + 1}'


@transaction.atomic
def gerar_excel(request, destino):
    wb = Workbook(write_only=True)
    # A tabela principal vem primeiro, com todos os vínculos lado a lado.
    with transaction.atomic():
        consolidado_por_estabelecimento(wb, request)
    guia = nova_aba(wb, 'Leia-me', ['Item', 'Descrição'])
    for item in [
        ('Gerado em', timezone.localtime().strftime('%d/%m/%Y %H:%M:%S')),
        ('Tabela principal', 'Consolidado por estabelecimento reúne dados cadastrais e todos os vínculos na mesma linha. Históricos extensos continuam nas linhas seguintes, com o mesmo ID e a coluna Parte. As demais abas são complementares.'),
        ('Abrangência', 'Todos os estabelecimentos, ativos e inativos, e todo o histórico. Filtros da tela não se aplicam.'),
        ('Vínculos', 'Número da pasta é a identificação usada no sistema. Pode repetir entre cadastros inativos e ativos. Para cruzar históricos sem misturar estabelecimentos, use estabelecimento_id com id da aba Estabelecimentos: são referências internas fixas.'),
        ('Numeração', 'As colunas de pasta mostram a numeração atual. O Histórico de pastas registra as alterações e a situação do cadastro; documentos emitidos preservam o conteúdo original.'),
        ('Usuários', 'Vínculos com usuários incluem ID e nome de exibição, sem credenciais.'),
        ('Situação do cadastro', 'No campo status de Estabelecimentos: falso = ativo; verdadeiro = inativo.'),
        ('Documentos', 'Os anexos não estão embutidos. Use o link de download com login no sistema.'),
        ('Conteúdo dos termos', 'Campos personalizados e snapshots estão na aba Conteúdo termos, identificados por aba, registro e estabelecimento.'),
        ('Textos extensos', 'Textos acima do limite de uma célula são divididos em partes numeradas na aba Textos extensos. Junte as partes em ordem.'),
        ('Sem vínculo', 'Denúncias sem estabelecimento também são incluídas, com estabelecimento_id vazio. Agenda e arquivos gerais não possuem vínculo cadastral.'),
        ('Datas', 'Datas e horas no fuso horário configurado no sistema.'),
    ]:
        linha(guia, item)
    conteudos = nova_aba(wb, 'Conteúdo termos', ['Aba', 'Registro ID', 'Estabelecimento ID interno', 'Número da pasta', 'Campo', 'Parte', 'Valor'])
    extensos = nova_aba(wb, 'Textos extensos', ['Aba', 'Registro ID', 'Campo', 'Parte', 'Texto'])
    with transaction.atomic():
        for nome, model in ABAS:
            campos = [f for f in model._meta.concrete_fields if not isinstance(f, models.JSONField)]
            titulos = []
            for campo in campos:
                titulos.append(campo.attname)
                if campo.is_relation:
                    titulos.append(f'{campo.name} — nome')
                    if campo.name == 'estabelecimento':
                        titulos.append('Número da pasta')
                if campo.choices:
                    titulos.append(f'{campo.name} — descrição')
            if model == Documento:
                titulos.append('Link para download')
            if model == NotificacaoEmitida:
                titulos.extend(['Número formatado', 'Link para documento'])
            aba = nova_aba(wb, nome, titulos)
            ordem = ('numero_pasta', 'pk') if model == Estabelecimento else ('estabelecimento__numero_pasta', 'estabelecimento_id', 'pk')
            registros = model.objects.select_related(*[f.name for f in campos if f.is_relation]).order_by(*ordem)
            total = 0
            for registro in registros.iterator(chunk_size=1000):
                valores = []
                for campo in campos:
                    valor = getattr(registro, campo.attname)
                    if isinstance(campo, models.FileField):
                        valor = valor.name
                    if isinstance(campo, models.UUIDField):
                        valor = str(valor)
                    if isinstance(valor, str) and len(valor) > 32767:
                        for indice, inicio in enumerate(range(0, len(valor), 32000), 1):
                            linha(extensos, [nome, registro.pk, campo.name, indice, valor[inicio:inicio + 32000]])
                        valor = 'Texto completo na aba Textos extensos'
                    valores.append(valor)
                    if campo.is_relation:
                        relacionado = getattr(registro, campo.name)
                        if relacionado is None:
                            valores.append('')
                        elif hasattr(relacionado, 'get_full_name'):
                            valores.append(relacionado.get_full_name() or relacionado.username)
                        else:
                            valores.append(str(relacionado))
                        if campo.name == 'estabelecimento':
                            valores.append(relacionado.numero_pasta if relacionado else None)
                    if campo.choices:
                        valores.append(getattr(registro, f'get_{campo.name}_display')())
                if model == Documento:
                    valores.append(request.build_absolute_uri(reverse('register_view:baixar_documento', args=[registro.pk])))
                if model == NotificacaoEmitida:
                    valores.extend([registro.numero_formatado, request.build_absolute_uri(reverse('termos:emitida', args=[registro.pk]))])
                linha(aba, valores)
                total += 1
                for campo in model._meta.concrete_fields:
                    if isinstance(campo, models.JSONField):
                        for chave, valor in getattr(registro, campo.name).items():
                            texto = valor if isinstance(valor, str) else json.dumps(valor, ensure_ascii=False)
                            for indice, inicio in enumerate(range(0, max(1, len(texto)), 32000), 1):
                                linha(conteudos, [nome, registro.pk, registro.estabelecimento_id, registro.estabelecimento.numero_pasta, chave, indice, texto[inicio:inicio + 32000]])
            aba.auto_filter.ref = f'A1:{get_column_letter(len(titulos))}{total + 1}'
            linha(guia, [nome, f'{total} registro(s)'])
    wb.save(destino)


@login_required
@require_GET
def exportar_consolidado(request):
    arquivo = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    try:
        gerar_excel(request, arquivo)
        arquivo.seek(0)
    except Exception:
        arquivo.close()
        raise
    response = FileResponse(arquivo, as_attachment=True,
        filename=f'consolidado_estabelecimentos_{timezone.localdate():%Y-%m-%d}.xlsx',
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Cache-Control'] = 'private, no-store'
    return response
